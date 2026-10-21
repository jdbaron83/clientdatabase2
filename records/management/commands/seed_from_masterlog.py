"""Load clients, episodes, authorizations, OCAs, MAT medications and the waitlist from the MasterLog workbook.

Sheets read: Master Log, MAT Outpatient Log (2), Waitlist, Waitlist Archive.
Everything else in the workbook (bed sheets, Outpatient Log, CBO, Sheet2, report tabs) is ignored.

Usage:
    python manage.py seed_from_masterlog "MasterLog.xlsx" [--reset] [--dry-run] [--review seed_review.csv]
"""

import csv
import re
from collections import Counter
from datetime import date, datetime, timedelta

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from records import audit, crypto
from records.models import (Authorization, Client, DischargeReason, Episode, FundingSource, LevelOfCare, Medication,
                            Oca, WaitlistEntry)

RECORD_ID = re.compile(r"^[A-Z]\d{2}-\d+(-\d+)?$")
NAME_SUFFIX = re.compile(r"\.\d+\s*$")
BED = re.compile(r"\bBed\s*#?\s*(\d+)\s*;?", re.I)
# OCA codes look like MS003, MSOTR, MSSM7, MH001, MS003CF2425; COSSAP is also an SFBHN OCA.
OCA_CODE = re.compile(r"^(M[SH][A-Z0-9]{3,}|COSSAP)$")
OCA_FUNDERS = ("SFBHN", "SFBHNMH")

FUNDING_ALIASES = {
    "self-pay": "Self-Pay",
    "mbpd": "MB PD",
    "mstrv": "Transitional Voucher - SFBHN MSTRV",
    "unfunded waitlist": "Unfunded",
}

DISCHARGE_REASONS = [
    (r"success.*(op|outpat|aftercare)|moved to aftercare", DischargeReason.SUCCESSFUL_TO_OP),
    (r"success", DischargeReason.SUCCESSFUL),
    (r"^asa", DischargeReason.ASA),
    (r"elope", DischargeReason.ELOPEMENT),
    (r"admin", DischargeReason.ADMIN),
    (r"apprehend", DischargeReason.APPREHENDED),
    (r"incarcer|jail", DischargeReason.INCARCERATED),
    (r"transfer", DischargeReason.TRANSFER),
    (r"hlc|hospital", DischargeReason.HOSPITAL),
    (r"readmit", DischargeReason.READMIT),
    (r"intake", DischargeReason.NO_INTAKE),
    (r"stopped reporting", DischargeReason.STOPPED_REPORTING),
]

EXCEL_EPOCH = datetime(1899, 12, 30)
STALE_BEFORE = date.today() - timedelta(days=730)


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    return "" if s in ("0", "#REF!", "#N/A", "#VALUE!", "None", "nan") else s


def to_date(value):
    """Parse the many date shapes in the workbook. Returns None for blanks, placeholders and 1900 filler dates."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        d = value.date()
    elif isinstance(value, date):
        d = value
    elif isinstance(value, (int, float)):
        if not 1000 < value < 80000:
            return None
        d = (EXCEL_EPOCH + timedelta(days=float(value))).date()
    else:
        s = str(value).strip()
        d = None
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%m-%d-%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                d = datetime.strptime(s, fmt).date()
                break
            except ValueError:
                continue
        if d is None:
            return None
    return d if 1920 <= d.year <= date.today().year + 2 else None


def to_int(value):
    try:
        n = int(float(value))
    except (TypeError, ValueError):
        return None
    return n if n >= 0 else None


def split_name(raw):
    name = NAME_SUFFIX.sub("", text(raw)).strip()
    name = re.sub(r"\s+", " ", name)
    if not name:
        return "", ""
    if "," in name:
        last, first = [p.strip() for p in name.split(",", 1)]
        return first, last
    parts = name.split(" ")
    if len(parts) == 1:
        return parts[0], ""
    return " ".join(parts[:-1]), parts[-1]


def discharge_reason(*values):
    for value in values:
        s = text(value).lower()
        if not s:
            continue
        for pattern, reason in DISCHARGE_REASONS:
            if re.search(pattern, s):
                return reason
        return DischargeReason.OTHER
    return ""


def level_for(name):
    n = name.lower()
    if re.search(r"\bmat\b", n):
        return LevelOfCare.MAT
    if re.search(r"outpatient|\bop\b|\bsor\b|aftercare|peer service", n):
        return LevelOfCare.OUTPATIENT
    if "unfunded" in n:
        return LevelOfCare.OTHER
    return LevelOfCare.RESIDENTIAL


class Seeder:
    def __init__(self, stdout):
        self.stdout = stdout
        self.stats = Counter()
        self.review = []
        self.funding = {}
        self.ocas = {}
        self.by_ssn = {}
        self.by_name_dob = {}
        self.by_name = {}
        self.episode_keys = {}

    # -- lookups -----------------------------------------------------------------------------------------------
    def funding_source(self, raw):
        name = re.sub(r"\s+", " ", text(raw)).strip(" ;,.")
        name = re.sub(r"^OMB\s*(\d)$", r"OMB \1", name, flags=re.I)
        if not name:
            return None
        name = FUNDING_ALIASES.get(name.lower(), name)
        key = name.lower()
        if key not in self.funding:
            fs, created = FundingSource.objects.get_or_create(name=name, defaults={"level_of_care": level_for(name)})
            self.stats["funding sources created"] += created
            self.funding[key] = fs
        return self.funding[key]

    def oca(self, raw):
        code = text(raw).upper().replace(" ", "")
        if not code:
            return None
        if not OCA_CODE.match(code):
            return None  # Older rows put a counselor's name or initials here; those are not OCAs.
        if code not in self.ocas:
            oca, created = Oca.objects.get_or_create(code=code)
            if created:
                oca.funding_sources.set([self.funding_source(n) for n in OCA_FUNDERS])
                self.stats["OCAs created"] += 1
            self.ocas[code] = oca
        return self.ocas[code]

    def client(self, name_raw, dob_raw, ssn_raw, sheet, row, **extra):
        first, last = split_name(name_raw)
        if not first:
            return None
        dob = to_date(dob_raw)
        ssn = crypto.normalize_ssn(ssn_raw)
        name_key = f"{first} {last}".lower()
        client = None
        if ssn:
            client = self.by_ssn.get(ssn)
        if client is None and dob:
            client = self.by_name_dob.get((name_key, dob))
        if client is None and not ssn and not dob:
            candidates = self.by_name.get(name_key, [])
            if len(candidates) == 1:
                client = candidates[0]
        if client is None:
            client = Client(first_name=first[:100], last_name=last[:100], date_of_birth=dob)
            client.set_ssn(ssn)
            for k, v in extra.items():
                setattr(client, k, v)
            client.save()
            self.stats["clients created"] += 1
            self.by_name.setdefault(name_key, []).append(client)
        else:
            changed = False
            if ssn and not client.ssn_hash:
                client.set_ssn(ssn)
                changed = True
            if dob and not client.date_of_birth:
                client.date_of_birth = dob
                changed = True
            for k, v in extra.items():
                if v and not getattr(client, k):
                    setattr(client, k, v)
                    changed = True
            if changed:
                client.save()
        if ssn:
            self.by_ssn[ssn] = client
        if dob:
            self.by_name_dob[(name_key, dob)] = client
        return client

    def flag(self, sheet, row, reason, detail=""):
        self.review.append({"sheet": sheet, "row": row, "reason": reason, "detail": detail})
        self.stats[f"review: {reason}"] += 1

    # -- sheets -----------------------------------------------------------------------------------------------
    def master_log(self, ws):
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        col = {}
        for i, h in enumerate(header):
            if h is not None:
                col.setdefault(str(h).strip(), i)

        def get(r, name):
            i = col.get(name)
            return r[i] if i is not None and i < len(r) else None

        for n, r in enumerate(rows, start=2):
            record_id = text(get(r, "Client ID"))
            if not RECORD_ID.match(record_id):
                continue  # filler and formula rows
            self.stats["master log rows"] += 1
            name = get(r, "Name")
            admit = to_date(get(r, "Admit Date"))
            funding = self.funding_source(get(r, "Funding Source"))
            if not admit or not funding:
                self.flag("Master Log", n, "missing admit date or funding source", f"{record_id} {text(name)}")
                continue
            client = self.client(name, get(r, "DOB"), get(r, "SSN"), "Master Log", n)
            if client is None:
                self.flag("Master Log", n, "missing name", record_id)
                continue
            discharge = to_date(get(r, "Discharge Date"))
            if discharge and discharge < admit:
                self.flag("Master Log", n, "discharge before admit; discharge left blank", record_id)
                discharge = None
            key = (client.pk, admit, funding.pk)
            if key in self.episode_keys:
                self.stats["duplicate episode rows skipped"] += 1
                continue
            benefits = text(get(r, "Disability/SSI/Unemployment"))
            bed_match = BED.search(benefits)
            bed = bed_match.group(1) if bed_match else ""
            if bed_match:
                benefits = BED.sub("", benefits).strip(" ;-")
            dc_note = text(get(r, "DC Note"))
            episode = Episode.objects.create(
                client=client,
                record_id=record_id,
                level_of_care=funding.level_of_care,
                funding_source=funding,
                initial_funding_source=self.funding_source(get(r, "Initial Funding Source")),
                secondary_funding=text(get(r, "Secondary Funding Source/Notes"))[:255],
                admit_date=admit,
                discharge_date=discharge,
                discharge_reason=discharge_reason(get(r, "D/C"), dc_note) if discharge else "",
                discharge_note=(dc_note or (text(get(r, "D/C")) if discharge else ""))[:255],
                bed=bed,
                diagnosis_code=text(get(r, "DxCode")) or (text(get(r, "Dx")) if re.match(r"^F\d", text(get(r, "Dx"))) else ""),
                benefits_notes=benefits[:255],
            )
            self.episode_keys[key] = episode
            self.stats["episodes created"] += 1
            if not discharge and admit < STALE_BEFORE:
                self.flag("Master Log", n, "still open with no discharge date after 2+ years",
                          f"{record_id} {funding} admitted {admit:%m/%d/%Y}")
            self.authorization(r, get, episode, funding, n)

    def authorization(self, r, get, episode, funding, n):
        auth_raw = text(get(r, "Authorization#"))
        auth_number = "" if auth_raw.lower().startswith("needauth") else auth_raw
        end = to_date(get(r, "Auth End Date"))
        days = to_int(get(r, "Auth Days"))
        oca_raw = text(get(r, "OCA"))
        oca = self.oca(oca_raw)
        if oca and funding.name not in OCA_FUNDERS:
            self.flag("Master Log", n, "OCA on a non-SFBHN funding source; OCA not loaded",
                      f"{episode.record_id} {funding} {oca}")
            oca = None
        if not (auth_raw or end or oca):
            return
        if end and days and end < episode.admit_date:
            end = None
        Authorization.objects.create(
            episode=episode, funding_source=funding, oca=oca, auth_number=auth_number[:60],
            start_date=episode.admit_date if (auth_number or end) else None, end_date=end,
            days_authorized=days or None,
        )
        self.stats["authorizations created"] += 1

    def mat_log(self, ws):
        rows = ws.iter_rows(values_only=True)
        next(rows)
        # Columns: 6 Client ID, 7 Name, 8 DOB, 9 age, 10 SSN, 11 Admit, 12 Discharge, 13 Funding, 14-16 demographics,
        # 35-40 MAT type, dosage, frequency, quantity, fill date, next fill date.
        for n, r in enumerate(rows, start=2):
            r = list(r) + [None] * 45
            record_id = text(r[6])
            if not RECORD_ID.match(record_id):
                continue
            self.stats["MAT log rows"] += 1
            admit = to_date(r[11])
            funding = self.funding_source(r[13]) or self.funding_source("SFBHN SOR")
            client = self.client(r[7], r[8], r[10], "MAT Outpatient Log (2)", n,
                                 gender=text(r[14]).capitalize(), race=text(r[15]), ethnicity=text(r[16]))
            if client is None or not admit:
                self.flag("MAT Outpatient Log (2)", n, "missing name or admit date", record_id)
                continue
            if funding.level_of_care == LevelOfCare.RESIDENTIAL:
                funding.level_of_care = LevelOfCare.MAT
                funding.save()
            key = (client.pk, admit, funding.pk)
            episode = self.episode_keys.get(key)
            if episode is None:
                discharge = to_date(r[12])
                episode = Episode.objects.create(
                    client=client, record_id=record_id, level_of_care=LevelOfCare.MAT, funding_source=funding,
                    admit_date=admit, discharge_date=discharge if discharge and discharge >= admit else None,
                    discharge_reason=DischargeReason.OTHER if discharge else "",
                )
                self.episode_keys[key] = episode
                self.stats["episodes created"] += 1
            drug = text(r[35])
            if drug and drug != "1":
                Medication.objects.create(
                    episode=episode, drug=drug[:100], dosage=text(r[36])[:50], frequency=text(r[37])[:50],
                    quantity=to_int(r[38]), fill_date=to_date(r[39]), next_fill_date=to_date(r[40]),
                )
                self.stats["medications created"] += 1

    def waitlist(self, ws):
        rows = list(ws.iter_rows(values_only=True))
        # Row 2 is the header: 1 Rank, 2 Name, 3 DOB, 4 SSN, 5 Admit, 6 Discharge, 7 Funding, 8 Secondary,
        # 9 Disability/notes, 10 Referral date, 11 Nota bene, 12 Wait (date added).
        for n, r in enumerate(rows[2:], start=3):
            r = list(r) + [None] * 16
            if not text(r[2]):
                continue
            self.stats["waitlist rows"] += 1
            client = self.client(r[2], r[3], r[4], "Waitlist", n)
            if client is None:
                continue
            on_hold = text(r[5]).lower() == "on hold"
            notes = "; ".join(p for p in [text(r[9]), text(r[11])] if p)
            WaitlistEntry.objects.create(
                client=client, funding_source=self.funding_source(r[7]),
                status=WaitlistEntry.Status.ON_HOLD if on_hold else WaitlistEntry.Status.WAITING,
                rank=to_int(r[1]), date_added=to_date(r[12]) or to_date(r[10]),
                referral_date=to_date(r[10]), secondary_funding=text(r[8])[:255], notes=notes,
            )
            self.stats["waitlist entries created"] += 1

    def waitlist_archive(self, ws):
        rows = ws.iter_rows(values_only=True)
        next(rows)
        # 7 Name, 8 DOB, 9 SSN, 10 Admit/status, 11 outcome note, 12 Funding, 13 Secondary, 14 notes, 17 Wait, 18 Archive date.
        for n, r in enumerate(rows, start=2):
            r = list(r) + [None] * 20
            if not text(r[7]):
                continue
            self.stats["waitlist archive rows"] += 1
            if to_date(r[8]) is None and text(r[8]) and not text(r[8]).isdigit():
                # A few rows have the name split across the first/last columns and everything shifted right.
                self.flag("Waitlist Archive", n, "row layout shifted; check by hand", f"{text(r[7])} {text(r[8])}")
                continue
            client = self.client(r[7], r[8], r[9], "Waitlist Archive", n)
            if client is None:
                continue
            admitted = to_date(r[10])
            archived = to_date(r[18]) or admitted
            WaitlistEntry.objects.create(
                client=client, funding_source=self.funding_source(r[12]),
                status=WaitlistEntry.Status.ADMITTED if admitted else WaitlistEntry.Status.DECLINED,
                date_added=to_date(r[17]),
                secondary_funding=text(r[13])[:255], notes=text(r[14]),
                closed_date=archived, closed_reason=text(r[11])[:255] or "Archived from spreadsheet",
            )
            self.stats["waitlist entries created"] += 1


class Command(BaseCommand):
    help = "Seed the database from the MasterLog workbook."

    def add_arguments(self, parser):
        parser.add_argument("workbook")
        parser.add_argument("--reset", action="store_true", help="Delete all client data before loading.")
        parser.add_argument("--dry-run", action="store_true", help="Load inside a transaction and roll it back.")
        parser.add_argument("--review", default="seed_review.csv",
                            help="Where to write rows that need a human look. Contains client names: keep it private.")

    def handle(self, workbook, reset, dry_run, review, **options):
        try:
            wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Workbook not found: {workbook}")
        seeder = Seeder(self.stdout)
        with audit.suspended(), transaction.atomic():
            if reset:
                for model in (WaitlistEntry, Medication, Authorization, Episode, Client, Oca, FundingSource):
                    model.objects.all().delete()
            elif Client.objects.exists():
                raise CommandError("The database already has clients. Use --reset to replace them.")
            for sheet, loader in [("Master Log", seeder.master_log), ("MAT Outpatient Log (2)", seeder.mat_log),
                                  ("Waitlist", seeder.waitlist), ("Waitlist Archive", seeder.waitlist_archive)]:
                if sheet in wb.sheetnames:
                    self.stdout.write(f"Loading {sheet}...")
                    loader(wb[sheet])
                else:
                    self.stderr.write(f"Sheet not found, skipped: {sheet}")
            if dry_run:
                transaction.set_rollback(True)

        for key, value in sorted(seeder.stats.items()):
            self.stdout.write(f"  {key}: {value}")
        if seeder.review:
            with open(review, "w", newline="") as fh:
                writer = csv.DictWriter(fh, fieldnames=["sheet", "row", "reason", "detail"])
                writer.writeheader()
                writer.writerows(seeder.review)
            self.stdout.write(f"{len(seeder.review)} rows need review: {review}")
        self.stdout.write(self.style.SUCCESS("Dry run, nothing saved." if dry_run else "Seed complete."))
